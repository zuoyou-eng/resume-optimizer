"""接口测试：按《系统设计文档》7.3 错误码表逐条设计用例。"""
from io import BytesIO

from app.config import settings


def _upload(client, filename: str, content: bytes):
    return client.post(
        "/api/v1/resumes",
        files={"file": (filename, BytesIO(content), "application/octet-stream")},
    )


class TestUploadApi:
    def test_upload_txt_success(self, client, upload_dir, sample_resume_text):
        resp = _upload(client, "resume.txt", sample_resume_text.encode("utf-8"))
        body = resp.json()
        assert body["code"] == 0
        data = body["data"]
        assert data["status"] == "parsed"
        assert data["structure"]["basic_info"]["email"] == "zhangsan@example.com"
        assert data["structure"]["education"], "教育背景应被解析出来"
        assert data["structure"]["experience"], "实习经历应被解析出来"
        # 统一响应四要素齐全
        assert body["trace_id"].startswith("req-")
        assert body["message"] == "上传并解析成功"

    def test_upload_docx_success(self, client, upload_dir):
        """Word 上传：用 python-docx 现场构造一个真实 docx。"""
        from docx import Document

        doc = Document()
        doc.add_paragraph("张三")
        doc.add_paragraph("邮箱：zhangsan@example.com")
        doc.add_paragraph("教育背景")
        doc.add_paragraph("浙江大学 计算机科学与技术 本科 2020.09-2024.06")
        buf = BytesIO()
        doc.save(buf)

        resp = _upload(client, "resume.docx", buf.getvalue())
        body = resp.json()
        assert body["code"] == 0
        assert body["data"]["structure"]["basic_info"]["email"] == "zhangsan@example.com"
        assert body["data"]["structure"]["education"][0]["school"] == "浙江大学"

    def test_upload_unsupported_type_1001(self, client, upload_dir):
        resp = _upload(client, "virus.exe", b"MZ\x90\x00 malicious")
        body = resp.json()
        assert body["code"] == 1001

    def test_upload_fake_extension_1001(self, client, upload_dir):
        """伪装扩展名：文件名是 .pdf，内容却不是 PDF 签名 → 必须拦截（NFR-05）。"""
        resp = _upload(client, "fake.pdf", b"this is plain text, not a pdf")
        body = resp.json()
        assert body["code"] == 1001

    def test_upload_too_large_1002(self, client, upload_dir, monkeypatch):
        monkeypatch.setattr(settings, "MAX_FILE_SIZE_MB", 0)
        resp = _upload(client, "big.txt", b"x" * 1024)
        body = resp.json()
        assert body["code"] == 1002

    def test_upload_empty_file_1003(self, client, upload_dir):
        resp = _upload(client, "empty.txt", b"")
        body = resp.json()
        assert body["code"] == 1003

    def test_upload_no_text_content_1003(self, client, upload_dir):
        """PDF 签名正确但文件损坏/无文本 → 1003 或 2001（解析失败）。"""
        resp = _upload(client, "blank.pdf", b"%PDF-1.4 minimal fake")
        body = resp.json()
        assert body["code"] in (1001, 1003, 2001)


class TestStructureApi:
    def test_get_structure_after_upload(self, client, upload_dir, sample_resume_text):
        resume_id = _upload(
            client, "resume.txt", sample_resume_text.encode("utf-8")
        ).json()["data"]["resume_id"]

        resp = client.get(f"/api/v1/resumes/{resume_id}/structure")
        body = resp.json()
        assert body["code"] == 0
        data = body["data"]
        assert data["parse_status"] == "success"
        assert data["rule_version"] == settings.RULE_VERSION
        assert len(data["trace"]) > 0
        # FR-03：未识别内容在接口输出中可见
        assert isinstance(data["unrecognized"], list)
        assert any("运动" in u["text"] for u in data["unrecognized"])

    def test_get_structure_not_found_2003(self, client, upload_dir):
        resp = client.get("/api/v1/resumes/nonexistent-id/structure")
        assert resp.json()["code"] == 2003


class TestDeleteApi:
    def test_delete_then_not_accessible(self, client, upload_dir, sample_resume_text):
        """NFR-05：删除后简历不可再访问。"""
        resume_id = _upload(
            client, "resume.txt", sample_resume_text.encode("utf-8")
        ).json()["data"]["resume_id"]

        resp = client.delete(f"/api/v1/resumes/{resume_id}")
        assert resp.json()["code"] == 0

        resp = client.get(f"/api/v1/resumes/{resume_id}/structure")
        assert resp.json()["code"] == 2003

    def test_delete_not_found_2003(self, client, upload_dir):
        resp = client.delete("/api/v1/resumes/nonexistent-id")
        assert resp.json()["code"] == 2003


class TestListApi:
    def test_list_contains_uploaded(self, client, upload_dir, sample_resume_text):
        _upload(client, "resume.txt", sample_resume_text.encode("utf-8"))
        resp = client.get("/api/v1/resumes")
        body = resp.json()
        assert body["code"] == 0
        assert len(body["data"]) >= 1
        item = body["data"][0]
        assert {"id", "original_name", "file_type", "status", "created_at"} <= set(item)
