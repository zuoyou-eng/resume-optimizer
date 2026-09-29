"""Word (.docx) 提取器：python-docx。

除正文段落外，也提取表格单元格文本——表格排版是简历的常见形式，
不提取表格会直接丢失内容（违反 FR-03 严禁静默丢失）。
"""
from app.errors import BizError, ErrorCode
from app.parser.base import ExtractedDocument, TextExtractor


class DocxExtractor(TextExtractor):
    name = "python-docx"

    def extract(self, file_path: str) -> ExtractedDocument:
        try:
            from docx import Document
        except ImportError as e:  # pragma: no cover
            raise BizError(
                ErrorCode.PARSE_FAILED, "Word 解析依赖未安装（python-docx）"
            ) from e

        paragraphs: list[str] = []
        try:
            doc = Document(file_path)
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    paragraphs.append(text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
        except Exception as e:
            raise BizError(ErrorCode.PARSE_FAILED, f"Word 解析失败：{e}") from e

        return ExtractedDocument(
            paragraphs=paragraphs, pages=1, source=self.name
        )
